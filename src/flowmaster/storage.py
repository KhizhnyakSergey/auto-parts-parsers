"""Persist scraped Flowmaster product rows to JSON Lines (raw/complete) and Excel.

Excel columns split into fixed (always present, one getter per column) and
dynamic groups whose *set* isn't known until we've seen every row: Specs
labels differ per product, and the number of photos/videos/instructions/notes
varies per product too.
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import FlowmasterRow

_EXCEL_CELL_LIMIT = 32000

_FIXED_COLUMNS_HEAD = [
    ("sku", lambda r: r.sku),
    ("title", lambda r: r.title),
    ("url", lambda r: r.url),
    ("breadcrumbs", lambda r: r.breadcrumbs),
    ("brand", lambda r: r.brand),
    ("short_description", lambda r: r.short_description),
    ("stock_status", lambda r: r.stock_status),
    ("in_stock", lambda r: r.in_stock),
    ("quantity", lambda r: r.quantity),
    ("price", lambda r: r.price),
    ("currency", lambda r: r.currency),
    ("lifecycle_status", lambda r: r.lifecycle_status),
    ("categories", lambda r: ", ".join(r.categories)),
]

_FIXED_COLUMNS_MID = [
    ("Overview", lambda r: r.overview_html),
]

_EMISSIONS_COLUMN = [
    ("Emissions", lambda r: r.emissions),
]

_INSTRUCTION_WARRANTY_COLUMN = [
    ("Instruction Warranty", lambda r: r.instruction_warranty or ""),
]

_FITMENT_COLUMNS = [
    ("auto", lambda r: r.auto),
    ("Submodel", lambda r: r.submodel),
    ("Engine", lambda r: r.engine),
    ("Body Type", lambda r: r.body_type),
    ("Transmission", lambda r: r.transmission),
]

_TRAILING_COLUMNS = [
    ("created_at", lambda r: r.created_at),
    ("updated_at", lambda r: r.updated_at),
]

_WIDTHS = {
    "title": 45,
    "url": 40,
    "breadcrumbs": 45,
    "short_description": 45,
    "categories": 30,
    "Overview": 60,
    "Emissions": 40,
    "auto": 50,
}


def _excel_safe(value):
    if isinstance(value, str) and len(value) > _EXCEL_CELL_LIMIT:
        return value[:_EXCEL_CELL_LIMIT] + " …[truncated]"
    return value


def _spec_columns(rows: list[FlowmasterRow]) -> list[str]:
    names: set[str] = set()
    for row in rows:
        names.update(row.specs.keys())
    return sorted(names)


def _image_columns(rows: list[FlowmasterRow]) -> list[str]:
    max_images = max((len(row.images) for row in rows), default=0)
    return [f"Image {i + 1}" for i in range(max_images)]


def _video_columns(rows: list[FlowmasterRow]) -> list[str]:
    max_videos = max((len(row.videos) for row in rows), default=0)
    return [f"Video {i + 1}" for i in range(max_videos)]


def _instruction_columns(rows: list[FlowmasterRow]) -> list[str]:
    max_links = max((len(row.instructions) for row in rows), default=0)
    return [f"Instruction {i + 1}" for i in range(max_links)]


def _note_columns(rows: list[FlowmasterRow]) -> list[str]:
    max_notes = max((len(row.notes) for row in rows), default=0)
    return [f"Note {i + 1}" for i in range(max_notes)]


def _build_columns(rows: list[FlowmasterRow]):
    image_names = _image_columns(rows)
    video_names = _video_columns(rows)
    spec_names = _spec_columns(rows)
    instruction_names = _instruction_columns(rows)
    note_names = _note_columns(rows)

    image_columns = [
        (name, (lambda r, i=i: r.images[i] if i < len(r.images) else ""))
        for i, name in enumerate(image_names)
    ]
    video_columns = [
        (name, (lambda r, i=i: r.videos[i] if i < len(r.videos) else ""))
        for i, name in enumerate(video_names)
    ]
    spec_columns = [(name, (lambda r, n=name: r.specs.get(n, ""))) for name in spec_names]
    instruction_columns = [
        (name, (lambda r, i=i: r.instructions[i] if i < len(r.instructions) else ""))
        for i, name in enumerate(instruction_names)
    ]
    note_columns = [
        (name, (lambda r, i=i: r.notes[i] if i < len(r.notes) else ""))
        for i, name in enumerate(note_names)
    ]

    return (
        _FIXED_COLUMNS_HEAD
        + image_columns
        + video_columns
        + _FIXED_COLUMNS_MID
        + spec_columns
        + _EMISSIONS_COLUMN
        + _INSTRUCTION_WARRANTY_COLUMN
        + instruction_columns
        + note_columns
        + _FITMENT_COLUMNS
        + _TRAILING_COLUMNS
    )


def write_jsonl(rows: list[FlowmasterRow], path: Path) -> None:
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(row.model_dump_json() + "\n")


def append_jsonl(row: FlowmasterRow, path: Path) -> None:
    with path.open("a", encoding="utf-8") as f:
        f.write(row.model_dump_json() + "\n")


def write_xlsx(rows: list[FlowmasterRow], path: Path) -> None:
    columns = _build_columns(rows)

    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = "Products"

    ws.append([name for name, _ in columns])
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)

    for row in rows:
        ws.append([_excel_safe(getter(row)) for _, getter in columns])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    for idx, (name, _) in enumerate(columns, start=1):
        default_width = 40 if name.startswith(("Image ", "Video ", "Instruction")) else 16
        ws.column_dimensions[get_column_letter(idx)].width = _WIDTHS.get(name, default_width)

    wb.save(path)


def load_jsonl(path: Path) -> list[FlowmasterRow]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8") as f:
        return [FlowmasterRow.model_validate(json.loads(line)) for line in f if line.strip()]
